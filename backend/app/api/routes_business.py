"""Business-only, explicitly mapped executive analytics. No production schema changes."""
from datetime import date, datetime
import pandas as pd
from openpyxl import load_workbook
from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session
from app.db.session import get_db
from app.db.models import Tenant, User, UploadedDocument
from app.security.rate_limit import check_ask_rate_limit, RateLimitExceeded
from app.security.auth import AuthContext, get_current_user, require_active_subscription
from app.billing.plans import normalize_account_type
from app.agents.business_metrics import compute_overview, answer_overview, MetricInputError

router = APIRouter(prefix="/business", tags=["business"])
MAX_ROWS = 50_000


def require_business(ctx: AuthContext = Depends(get_current_user), db: Session = Depends(get_db)):
    tenant = db.query(Tenant).filter_by(id=ctx.tenant_id).first()
    if not tenant or normalize_account_type(tenant.account_type) != "business":
        raise HTTPException(403, "Executive Overview is available to business accounts.")
    user = db.query(User).filter_by(id=ctx.user_id, tenant_id=ctx.tenant_id).first()
    if not user or not {"querying", "document_retrieval"}.issubset(set(user.capabilities or [])):
        raise HTTPException(403, "Querying and document access are required.")
    return ctx


def load_source(db, ctx, document_id, sheet=0):
    doc = db.query(UploadedDocument).filter_by(id=document_id, tenant_id=ctx.tenant_id).first()
    if not doc:
        raise HTTPException(404, "Source not found.")
    if doc.kind not in ("csv", "xlsx"):
        raise HTTPException(422, "Executive Overview currently requires a CSV or Excel table.")
    try:
        if doc.kind == "csv":
            header = pd.read_csv(doc.file_path, nrows=0, skip_blank_lines=False)
            if len(header.columns) > 120:
                raise HTTPException(422, "Use a table with at most 120 columns.")
            frame = pd.read_csv(doc.file_path, nrows=MAX_ROWS + 1, dtype=str, skip_blank_lines=False)
        else:
            workbook = load_workbook(doc.file_path, read_only=True, data_only=True)
            try:
                columns = workbook.worksheets[sheet].max_column or 0
                if columns > 120:
                    raise HTTPException(422, "Use a table with at most 120 columns.")
            finally:
                workbook.close()
            frame = pd.read_excel(doc.file_path, sheet_name=sheet, nrows=MAX_ROWS + 1,
                                  usecols=range(min(columns, 120)), engine="openpyxl")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, "The selected sheet could not be read as a table.")
    if len(frame) > MAX_ROWS or len(frame.columns) > 120:
        raise HTTPException(422, "Use a flat table with at most 50,000 rows and 120 columns.")
    if frame.empty:
        raise HTTPException(422, "This source contains no data rows.")
    frame.columns = frame.columns.astype(str)
    user = db.query(User).filter_by(id=ctx.user_id).first()
    for col, values in (user.row_scope or {}).items():
        if col not in frame.columns:
            raise HTTPException(403, "This upload cannot enforce your row permissions.")
        frame = frame[frame[col].astype(str).isin([str(v) for v in values])]
    return doc, frame


@router.get("/sources/{document_id}")
def source_columns(document_id: str, sheet: int = 0, db: Session = Depends(get_db),
                   ctx: AuthContext = Depends(require_business), _paid=Depends(require_active_subscription)):
    doc, frame = load_source(db, ctx, document_id, sheet)
    return {"id": doc.id, "filename": doc.filename, "columns": list(frame.columns),
            "row_count": len(frame), "uploaded_at": doc.created_at.isoformat(), "version": doc.content_sha256}


class OverviewRequest(BaseModel):
    document_id: str
    sheet: int = Field(default=0, ge=0, le=99)
    date_column: str
    revenue_column: str | None = None
    expense_column: str | None = None
    currency_column: str | None = None
    group_column: str | None = None
    group_value: str | None = Field(default=None, max_length=500)
    rank_by: str = Field(default="revenue", pattern=r"^(revenue|expenses)$")
    currency: str = Field(pattern=r"^[A-Za-z]{3}$")
    include_duplicates: bool = False
    start: date
    end: date
    detail_date: date | None = None
    question: str | None = Field(default=None, max_length=1000)


@router.post("/overview")
def overview(body: OverviewRequest, db: Session = Depends(get_db),
             ctx: AuthContext = Depends(require_business), _paid=Depends(require_active_subscription)):
    try:
        check_ask_rate_limit("business-overview:" + ctx.user_id)
    except RateLimitExceeded:
        raise HTTPException(429, "Too many overview requests. Please try again shortly.")
    doc, frame = load_source(db, ctx, body.document_id, body.sheet)
    try:
        result = compute_overview(frame, body.model_dump(), body.start, body.end, body.detail_date)
    except MetricInputError as exc:
        raise HTTPException(422, str(exc)) from exc
    result["source"] = {"id": doc.id, "filename": doc.filename, "version": doc.content_sha256,
                         "uploaded_at": doc.created_at.isoformat(), "sheet": body.sheet}
    result["generated_at"] = datetime.utcnow().isoformat() + "Z"
    if body.question:
        result["answer"] = answer_overview(body.question, result)
    return result

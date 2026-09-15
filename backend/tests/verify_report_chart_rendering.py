"""
Regression guard for a real production report: a document-only analysis's
charts (line + bar) rendered fine on screen but the downloaded PDF report
never actually drew a chart at all - only a plain data table - and the
existing tests (verify_document_analysis_v2.py's checks 6/7) only asserted
"the file exists and has nonzero size", which a table-only PDF/PPTX also
satisfies. This file actually opens the generated PDF/PPTX back up and
checks for a real drawn chart / native chart object, not just a file on
disk, so a future regression back to "table instead of chart" gets caught.

Run from backend/:  PYTHONPATH=$(pwd) python tests/verify_report_chart_rendering.py
"""
import base64, os, tempfile

_tmp = tempfile.mkdtemp()
os.environ.setdefault("APP_SECRET_KEY", base64.urlsafe_b64encode(b"0" * 32).decode())
os.environ["METADATA_DB_URL"] = f"sqlite:///{os.path.join(_tmp, 'meta.db')}"
os.environ["ARTIFACTS_DIR"] = os.path.join(_tmp, "artifacts")
os.environ["DOCUMENTS_DIR"] = os.path.join(_tmp, "documents")

import pymupdf  # used here only to open the generated PDF back up and
                # inspect its actual vector drawing ops, never for
                # anything the app itself ships
from pptx import Presentation

from app.agents.report_generator import generate_report_pdf, _safe
from app.agents.presentation_generator import generate_presentation_pptx

_BASE_KWARGS = dict(
    question="q", metrics={},
    data_quality={"row_count": 0, "completeness_pct": 100.0, "duplicate_pct": 0.0, "notes": []},
    anomalies=[],
)

# A non-empty extraction_summary matters: {} is falsy, so an empty dict
# here would silently fall through to a different (and, as it turns out,
# separately buggy - an empty-string bullet crashes presentation_generator's
# _style_run) branch than the one real v2 document-only results use.
_INSIGHT = {
    "what": "x", "confidence": "high", "confidence_explanation": "x", "data_quality_caveat": "x",
    "next_question": "x",
    "extraction_summary": {"total_rows_or_items": 6, "sheets_or_pages_or_slides": 1, "extraction_confidence": "high", "flags": []},
    "key_findings": [{"finding": "x", "location": "x", "confidence": "high"}],
    "flagged_items": [],
}


def _pdf_drawings(path: str) -> list:
    doc = pymupdf.open(path)
    drawings = []
    for page in doc:
        drawings.extend(page.get_drawings())
    return drawings


# --- 1. a "line" chart draws real connected line segments in the PDF, not
#        a table - the exact shape (Spend as % of Sales by Month) from the
#        real report that motivated this fix ---
line_chart = [{
    "chart_type": "line", "title": "Spend as % of Sales by Month",
    "labels": ["Jan", "Feb", "Mar", "Apr", "May", "Jun"],
    "values": [40.2, 38.2, 38.2, 37.3, 36.3, 35.8],
    "unit": "%", "insight": "Declining spend ratio.", "location": "Operating Spend table",
}]
pdf_path1 = generate_report_pdf(
    title="Line chart test", insight=_INSIGHT,
    by_group=None, sql="-- n/a", query_id="q-line-1", charts=line_chart, **_BASE_KWARGS,
)
line_lines = [d for d in _pdf_drawings(pdf_path1) if any(item[0] == "l" for item in d.get("items", []))]
assert len(line_lines) >= 5, f"expected at least 5 connecting line segments for a 6-point line chart, got {len(line_lines)}"
print("1. OK  a 'line' chart draws real connected line segments in the PDF (not a data table)")

# --- 2. a "bar" chart draws real filled rectangles, one per category ---
bar_chart = [{
    "chart_type": "bar", "title": "December Operating Spend Breakdown by Category",
    "labels": ["Fuel", "Fleet", "Staff", "Warehouse", "Other"],
    "values": [8.4, 6.0, 6.2, 2.9, 2.5],
    "unit": "₦m", "insight": "Fuel leads.", "location": "table",
}]
pdf_path2 = generate_report_pdf(
    title="Bar chart test", insight=_INSIGHT,
    by_group=None, sql="-- n/a", query_id="q-bar-1", charts=bar_chart, **_BASE_KWARGS,
)
rects = [d for d in _pdf_drawings(pdf_path2) if d.get("fill") is not None and any(item[0] == "re" for item in d.get("items", []))]
assert len(rects) >= 5, f"expected at least 5 filled bar rectangles, got {len(rects)}"
print("2. OK  a 'bar' chart draws real filled rectangles in the PDF, one per category")

# --- 3. the OLD-shape DB path (by_group, no charts) ALSO gets a real drawn
#        bar chart now, not just the new v2 document-only shape ---
pdf_path3 = generate_report_pdf(
    title="by_group chart test", insight={"what": "x", "where": "N/A", "when": "N/A", "contributors": "N/A"},
    by_group=[{"group": "North", "total": 15.0}, {"group": "South", "total": 12.0}, {"group": "East", "total": 9.0}],
    sql="-- n/a", query_id="q-bg-1", **_BASE_KWARGS,
)
bg_rects = [d for d in _pdf_drawings(pdf_path3) if d.get("fill") is not None and any(item[0] == "re" for item in d.get("items", []))]
assert len(bg_rects) >= 3, f"expected at least 3 filled bars for by_group's 3 groups, got {len(bg_rects)}"
print("3. OK  the OLD-shape (by_group) DB path also draws a real bar chart, not a table")

# --- 4. degenerate chart data (empty, all-zero, single flat value) never
#        raises - these divide-by-zero-shaped inputs are exactly what a
#        real "no variation this month" or malformed model chart call
#        could produce ---
for degenerate in (
    [{"chart_type": "line", "title": "flat", "labels": ["A", "B", "C"], "values": [5, 5, 5], "unit": None, "insight": None, "location": None}],
    [{"chart_type": "bar", "title": "empty", "labels": [], "values": [], "unit": None, "insight": None, "location": None}],
    [{"chart_type": "bar", "title": "zeros", "labels": ["A", "B"], "values": [0, 0], "unit": None, "insight": None, "location": None}],
):
    p = generate_report_pdf(
        title="Degenerate test", insight=_INSIGHT,
        by_group=None, sql="-- n/a", query_id="q-degen", charts=degenerate, **_BASE_KWARGS,
    )
    assert os.path.exists(p) and os.path.getsize(p) > 0
print("4. OK  degenerate chart data (flat series, empty, all-zero) never raises")

# --- 4b. the Naira sign no longer becomes a bare "?" - a real exported
#         report showed "December Operating Spend Breakdown by Category
#         (?m)" because fpdf2's core fonts are Latin-1 and have no "₦"
#         glyph; _safe() now substitutes "NGN " instead of falling through
#         to its generic '?' replacement-char behavior. ---
assert _safe("Category (₦m)") == "Category (NGN m)", _safe("Category (₦m)")
assert _safe("₦80.7m") == "NGN 80.7m", _safe("₦80.7m")

naira_chart = [{
    "chart_type": "bar", "title": "December Operating Spend Breakdown by Category (₦m)",
    "labels": ["Fuel", "Fleet"], "values": [80.7, 63.3], "unit": "₦m",
    "insight": "Fuel leads at ₦80.7m.", "location": "table",
}]
naira_pdf = generate_report_pdf(
    title="Naira test", insight=_INSIGHT, by_group=None, sql="-- n/a",
    query_id="q-naira", charts=naira_chart, **_BASE_KWARGS,
)
naira_text = "".join(page.get_text() for page in pymupdf.open(naira_pdf))
assert "NGN" in naira_text, naira_text
assert "?" not in naira_text, naira_text
print("4b. OK  the Naira sign (₦) renders as 'NGN' in the PDF instead of a bare '?'")

# --- 5. PPTX: a "line" chart produces a real native LINE_MARKERS chart
#        object with the right categories and values, not a table ---
pptx_path1 = generate_presentation_pptx(
    title="Line chart test", insight=_INSIGHT,
    by_group=None, query_id="q-line-2", charts=line_chart, **_BASE_KWARGS,
)
prs1 = Presentation(pptx_path1)
chart_shapes1 = [s for slide in prs1.slides for s in slide.shapes if s.has_chart]
assert len(chart_shapes1) == 1, f"expected exactly 1 chart slide, got {len(chart_shapes1)}"
chart1 = chart_shapes1[0].chart
assert str(chart1.chart_type).startswith("LINE_MARKERS"), chart1.chart_type
assert list(chart1.plots[0].categories) == ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]
assert list(chart1.plots[0].series[0].values) == [40.2, 38.2, 38.2, 37.3, 36.3, 35.8]
print("5. OK  PPTX 'line' chart is a real native LINE_MARKERS chart object with the correct data")

# --- 6. PPTX: a "bar" chart produces a real COLUMN_CLUSTERED chart, and
#        by_group (the OLD-shape DB path) gets one too ---
pptx_path2 = generate_presentation_pptx(
    title="Bar chart test", insight=_INSIGHT,
    by_group=None, query_id="q-bar-2", charts=bar_chart, **_BASE_KWARGS,
)
prs2 = Presentation(pptx_path2)
chart2 = next(s.chart for slide in prs2.slides for s in slide.shapes if s.has_chart)
assert str(chart2.chart_type).startswith("COLUMN_CLUSTERED"), chart2.chart_type
assert list(chart2.plots[0].series[0].values) == [8.4, 6.0, 6.2, 2.9, 2.5]

pptx_path3 = generate_presentation_pptx(
    title="by_group chart test", insight={"what": "x", "where": "N/A", "when": "N/A", "contributors": "N/A"},
    by_group=[{"group": "North", "total": 15.0}, {"group": "South", "total": 12.0}],
    query_id="q-bg-2", **_BASE_KWARGS,
)
prs3 = Presentation(pptx_path3)
chart3_shapes = [s for slide in prs3.slides for s in slide.shapes if s.has_chart]
assert len(chart3_shapes) == 1
assert list(chart3_shapes[0].chart.plots[0].categories) == ["North", "South"]
print("6. OK  PPTX 'bar' chart and the OLD-shape by_group path both produce real native chart objects")

print("\nall report/presentation chart-rendering checks passed")

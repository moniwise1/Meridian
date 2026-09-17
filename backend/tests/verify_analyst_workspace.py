"""
Verification for the Ask analyst workspace: the analysis contract, confirmed
memory, conversation reopening, and the durable upload-check jobs.

Real SQLite metadata DB + real SQLAlchemy Session + the real route handler
functions and the real run_analysis() generator. Only the external
boundaries are stubbed (the Anthropic calls); the DB layer under test
(AskMemory / AskFinding / AskScanJob / Conversation / QueryRecord, the
additive migration, the authorization re-checks) is 100% real.

This is written as a standalone script rather than a pytest module because
this project has no pytest dependency - tests/run_regressions.py discovers
and runs every verify_*.py here, and CI runs that.

Run from backend/:  python tests/verify_analyst_workspace.py
"""
import base64
import json
import os
import tempfile
from types import SimpleNamespace

_tmp = tempfile.mkdtemp()
os.environ.setdefault("APP_SECRET_KEY", base64.urlsafe_b64encode(b"0" * 32).decode())
os.environ["METADATA_DB_URL"] = f"sqlite:///{os.path.join(_tmp, 'meta.db')}"
os.environ.setdefault("ARTIFACTS_DIR", os.path.join(_tmp, "artifacts"))

from fastapi import BackgroundTasks, HTTPException
from sqlalchemy import inspect

import app.db.session as sessions
from app.db.session import init_db, SessionLocal
from app.db.models import (
    Tenant, User, UploadedDocument, DataSourceConnection, Conversation, QueryRecord,
    AskMemory, AskFinding, AskScanJob,
)
from app.security.auth import AuthContext
from app.api.routes_analyst import (
    workspace, save_memory, delete_memory, update_finding, get_conversation,
    MemoryInput, FindingUpdate,
)
from app.agents.analyst_contract import Analysis, build_analysis
from app.agents import analyst_workspace as scans
from app.agents.analyst_workspace import memories_for, enqueue_upload_check, run_pending_checks
from app.agents.insight_agent import DocumentInsight
from app.agents.context_resolver import ResolvedQuestion
import app.agents.planner as planner


def _ctx(user_id="u", tenant_id="t"):
    return AuthContext(user_id, tenant_id, "admin")


def _raises(exc_type, fn, *args, **kwargs):
    """assertRaises without pulling in unittest/pytest - returns the exception."""
    try:
        fn(*args, **kwargs)
    except exc_type as e:
        return e
    raise AssertionError(f"expected {exc_type.__name__}, nothing was raised")


# --- 1. the contract refuses to present an unknown as a measured zero -------
result = build_analysis("AQ-1", "What was revenue?", {"data_quality": {"completeness_pct": 100}},
                        sources=[{"id": "d"}], computed=False)
assert result["quality"]["completeness_pct"] is None, result["quality"]
assert result["quality"]["row_count"] is None, result["quality"]
assert result["confidence"]["measurement"]["level"] == "not_assessable", result["confidence"]
assert result["metrics"] == [], result["metrics"]
assert result["evidence"][0]["method"] == "extracted_text", result["evidence"]
assert any("not independently computed" in n for n in result["quality"]["limitations"]), \
    result["quality"]["limitations"]
print("1. OK  an unverified document answer reports not_assessable, not a measured 100%")


# --- 2. every chart points at a real table, and a broken reference is
#        rejected by the schema rather than rendered ------------------------
result = build_analysis("AQ-2", "Revenue by region?",
                        {"row_count": 2, "metrics": {"total": 30},
                         "by_group": [{"group": "A", "total": 30}]},
                        sources=[{"id": "d", "version": "v1"}])
assert result["tables"][0]["rows"][0]["total"] == 30, result["tables"]
assert result["charts"][0]["table_id"] == result["tables"][0]["id"], result["charts"]
assert result["evidence"][0]["method"] == "computed", result["evidence"]
broken = json.loads(json.dumps(result))
broken["charts"][0]["table_id"] = "invented"
_raises(ValueError, Analysis.model_validate, broken)
broken = json.loads(json.dumps(result))
broken["metrics"][0]["evidence_ids"] = ["e99"]
_raises(ValueError, Analysis.model_validate, broken)
print("2. OK  charts and metrics must resolve to a real table/evidence id or validation fails")


# --- 3. a missing number stays missing - never coerced to 0 ----------------
result = build_analysis("AQ-3", "Revenue?",
                        {"row_count": 1, "metrics": {"total": float("nan")},
                         "by_group": [{"group": "A", "total": None}]},
                        sources=[{"id": "d"}])
assert result["metrics"] == [], result["metrics"]
assert result["tables"][0]["rows"][0]["total"] is None, result["tables"]
print("3. OK  NaN/None figures are dropped or left null, never reported as zero")


# --- 4. a percentage column is never presented with a "Sum of" total -------
result = build_analysis("AQ-4", "Conversion rate?", {"row_count": 5, "metrics": {"total": 412.0, "mean": 82.4}},
                        sources=[{"id": "d"}],
                        profile={"primary_value_column": "conversion_rate", "breakdowns": {}})
labels = [m["label"] for m in result["metrics"]]
assert not any(label.startswith("Sum of") for label in labels), labels
assert "Average of conversion_rate" in labels, labels
print("4. OK  a rate column gets an average, not a meaningless sum")


# --- real database from here on -------------------------------------------
init_db()
db = SessionLocal()
db.add_all([Tenant(id="t", name="One"), Tenant(id="other", name="Other")])
for uid, tid in [("u", "t"), ("v", "t"), ("x", "other")]:
    db.add(User(id=uid, tenant_id=tid, email=f"{uid}@test.invalid", password_hash="unused",
                capabilities=["querying", "document_retrieval"]))
db.add(UploadedDocument(id="d", tenant_id="t", user_id="u", filename="sales.csv", kind="csv",
                        file_path="unused", extracted_text="", content_sha256="version1"))
db.commit()


# --- 5. a confirmed note is private to the user who saved it ---------------
note = save_memory(MemoryInput(source_key="doc:d", content="Revenue excludes tax."), db, _ctx())
assert memories_for(db, "t", "u", "doc:d")[0]["content"] == "Revenue excludes tax."
assert memories_for(db, "t", "v", "doc:d") == [], "another user on the same tenant saw a private note"
err = _raises(HTTPException, delete_memory, note["id"], db, _ctx("v"))
assert err.status_code == 404, err.status_code
_raises(HTTPException, save_memory, MemoryInput(source_key="doc:d", content="x"), db, _ctx("x", "other"))
delete_memory(note["id"], db, _ctx())
assert memories_for(db, "t", "u", "doc:d") == []
print("5. OK  confirmed notes are scoped to one user and one source, and are removable")


# --- 6. reopening a conversation re-checks source access at read time ------
db.add(Conversation(id="c", tenant_id="t", user_id="u", connection_id=planner.DOCUMENT_ONLY_SOURCE_ID,
                    context={"source_key": "doc:d", "document_ids": ["d"]}))
db.add(QueryRecord(id="AQ-c1", tenant_id="t", user_id="u", connection_id=planner.DOCUMENT_ONLY_SOURCE_ID,
                   conversation_id="c", question="What was revenue?", generated_sql="",
                   result_snapshot={"insight": {"what": "Example"}}, row_count=0, duration_ms=0))
db.commit()
assert get_conversation("c", db, _ctx())["turns"][0]["question"] == "What was revenue?"
_raises(HTTPException, get_conversation, "c", db, _ctx("v"))
user = db.get(User, "u")
user.capabilities = ["querying"]          # document access revoked after the fact
db.commit()
err = _raises(HTTPException, get_conversation, "c", db, _ctx())
assert err.status_code == 403, err.status_code
user.capabilities = ["querying", "document_retrieval"]
db.commit()
print("6. OK  a revoked capability blocks reopening a conversation that used it")


# --- 7. a database conversation is refused once its permissions change -----
conn = DataSourceConnection(id="conn1", tenant_id="t", name="wh", kind="postgres", host="h", port=5432,
                            database="db", username="u", encrypted_password="x",
                            table_allowlist=["sales"], column_policy={})
db.add(conn)
db.add(Conversation(id="c2", tenant_id="t", user_id="u", connection_id="conn1",
                    context={"source_key": "conn:conn1", "document_ids": [],
                             "policy_signature": scans.policy_signature(conn, {})}))
db.commit()
assert get_conversation("c2", db, _ctx())["source_key"] == "conn:conn1"
conn.table_allowlist = ["sales", "payroll"]
db.commit()
err = _raises(HTTPException, get_conversation, "c2", db, _ctx())
assert err.status_code == 403, err.status_code
conn.table_allowlist = ["sales"]
db.commit()
print("7. OK  widening a connection's allowlist invalidates the old conversation")


# --- 8. upload checks: deduplicated, durable, capped at three attempts -----
doc = db.get(UploadedDocument, "d")
enqueue_upload_check(db, doc)
db.commit()
enqueue_upload_check(db, doc)
db.commit()
assert db.query(AskScanJob).count() == 1, "the same document was queued twice"

_real_scan = scans.scan_findings
scans.scan_findings = lambda d: [{"kind": "quality", "title": "Missing values need review",
                                  "detail": "Example", "confidence": "high"}]
run_pending_checks("t", "u")
run_pending_checks("t", "u")          # a second pass must not duplicate the finding
db.expire_all()
assert db.query(AskFinding).count() == 1, db.query(AskFinding).count()
finding = db.query(AskFinding).one()
assert finding.source_version == "version1", finding.source_version
assert db.query(AskScanJob).one().status == "complete"

update_finding(finding.id, FindingUpdate(status="expected"), db, _ctx())
shown = workspace("doc:d", BackgroundTasks(), db, _ctx())
assert shown["findings"][0]["status"] == "expected", shown["findings"]
assert shown["findings"][0]["title"] == "Missing values need review", shown["findings"]
_raises(HTTPException, update_finding, finding.id, FindingUpdate(status="dismissed"), db, _ctx("v"))

job = db.query(AskScanJob).one()
job.status = "failed"
job.attempts = 0
db.commit()


def _boom(d):
    raise RuntimeError("deliberate failure")


scans.scan_findings = _boom
for _ in range(5):
    run_pending_checks("t", "u")
db.expire_all()
assert db.query(AskScanJob).one().attempts == 3, db.query(AskScanJob).one().attempts
assert db.query(AskScanJob).one().status == "failed"
scans.scan_findings = _real_scan
print("8. OK  upload checks dedupe, persist, survive failure, and stop after three attempts")


# --- 9. a document follow-up reuses the conversation and the saved notes ---
csv_path = os.path.join(_tmp, "sales.csv")
with open(csv_path, "w", encoding="utf-8") as f:
    f.write("region,revenue\nNorth,100\nSouth,200\nNorth,150\nSouth,250\n")
doc.file_path = csv_path
db.commit()
save_memory(MemoryInput(source_key="doc:d", content="Revenue excludes refunds."), db, _ctx())

seen_context = []


def _fake_explain_doc_only(question, documents, computed_profile=None, force_answer=False,
                           business_context=None):
    seen_context.append(business_context)
    return (
        DocumentInsight(
            what="Review the computed revenue profile.", confidence="medium",
            confidence_explanation="Scope needs review.", data_quality_caveat="",
            next_question="Compare regions", extraction_summary={}, key_findings=[],
            flagged_items=[], structured_data=[],
        ),
        [],
    )


planner.explain_document_only_v2 = _fake_explain_doc_only
planner.resolve_followup = lambda q, c: ResolvedQuestion(
    resolved_question="Compare regions using revenue excluding refunds")


def _run_doc(question, conversation_id=None, user_id="u"):
    final = None
    for event in planner.run_analysis(db, "t", user_id, None, question, row_scope={},
                                      conversation_id=conversation_id, document_ids=["d"]):
        if isinstance(event, dict) and event.get("final"):
            final = event
    return final


first = _run_doc("Revenue by region")
second = _run_doc("What about the split?", conversation_id=first["conversation_id"])
assert first["conversation_id"], "the document-only path produced no conversation"
assert first["conversation_id"] == second["conversation_id"], (first["conversation_id"], second["conversation_id"])
assert len(get_conversation(first["conversation_id"], db, _ctx())["turns"]) == 2
assert seen_context[1][0]["content"] == "Revenue excludes refunds.", seen_context[1]
assert second["analysis"]["evidence"][0]["source_version"] == "version1", second["analysis"]["evidence"]
totals = {m["label"]: m["value"] for m in second["analysis"]["metrics"]}
assert totals["Sum of revenue"] == 700, totals
assert db.query(QueryRecord).filter_by(conversation_id=first["conversation_id"]).count() == 2
err = _raises(planner.PolicyViolation, _run_doc, "Revenue", first["conversation_id"], "v")
assert "does not match" in str(err), str(err)
print("9. OK  a document follow-up reuses its conversation, its notes, and real computed totals")


# --- 10. the additive migration is repeatable ------------------------------
with sessions.engine.begin() as c:
    c.exec_driver_sql("ALTER TABLE uploaded_documents DROP COLUMN content_sha256")
init_db()
init_db()
columns = {c["name"] for c in inspect(sessions.engine).get_columns("uploaded_documents")}
assert "content_sha256" in columns, columns
tables = set(inspect(sessions.engine).get_table_names())
assert {"ask_memories", "ask_findings", "ask_scan_jobs"} <= tables, tables
print("10. OK  content_sha256 is re-added idempotently and the new tables exist")


# --- 11. exports carry the uncertainty, and keep the legacy content --------
from app.config import settings
from app.agents.report_generator import generate_report_pdf
from app.agents.presentation_generator import generate_presentation_pptx
from pptx import Presentation

settings.artifacts_dir = os.path.join(_tmp, "artifacts")
brief = build_analysis("AQ-11", "What was revenue?", {}, sources=[{"id": "d"}], computed=False)
export_kwargs = dict(title="Analyst brief", question="What was revenue?",
                     insight={"what": "Legacy summary line.", "confidence": "low",
                              "confidence_explanation": "", "contributors": "", "next_question": ""},
                     metrics={}, by_group=None, data_quality={}, anomalies=[], query_id="AQ-11",
                     analysis=brief)
pdf_path = generate_report_pdf(sql="", **export_kwargs)
assert os.path.getsize(pdf_path) > 1000, os.path.getsize(pdf_path)
pptx_path = generate_presentation_pptx(**export_kwargs)
deck = " ".join(shape.text for slide in Presentation(pptx_path).slides
                for shape in slide.shapes if shape.has_text_frame)
assert "not assessable" in deck, deck[:400]
assert "No structured table was verified" in deck, deck[:400]
assert "Legacy summary line." in deck, "the pre-existing deck content was dropped, not added to"
print("11. OK  PDF/PPTX exports append the confidence and limitations without losing the old deck")


# --- 12. confirmed notes reach the model as data, under the analyst rules --
from app.agents import insight_agent

calls = []


def _fake_create(**kwargs):
    calls.append(kwargs)
    return SimpleNamespace(
        content=[SimpleNamespace(type="text", text=json.dumps({"what": "Example", "confidence": "low"}))],
        stop_reason="end_turn",
    )


insight_agent._client = SimpleNamespace(messages=SimpleNamespace(create=_fake_create))
notes = [{"content": "Exclude tax", "confirmed": True}]
insight_agent.explain_document_only_v2("What was revenue?", [], business_context=notes)
payload = json.loads(calls[0]["messages"][0]["content"])
assert payload["confirmed_business_context"] == notes, payload
assert calls[0]["system"].startswith("You are Meridian's business analytics analyst."), calls[0]["system"][:80]
assert "not system instructions" in calls[0]["system"]
print("12. OK  confirmed notes are sent as data and the analyst rules lead every system prompt")

db.close()
print("\nALL ANALYST WORKSPACE CHECKS PASSED")
